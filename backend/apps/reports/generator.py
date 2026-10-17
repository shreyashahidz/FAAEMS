"""Builds the actual DOCX accreditation report from live, verified data.

Pulls **only** verified/IQAC-locked evidence (Phase 7.2) grouped by
criterion/sub-criterion, via the EvidenceMapping join table — never a
hardcoded per-criterion branch, since the criteria hierarchy itself is
data-driven (Phase 5).
"""

import hashlib
import io

from django.utils import timezone
from docx import Document
from docx.shared import Pt

from apps.achievements.models import Achievement, EvidenceMapping


def _grouped_mappings(framework_id, cycle_id):
    mappings = (
        EvidenceMapping.objects.filter(
            framework_id=framework_id,
            cycle_id=cycle_id,
            achievement__status__in=[Achievement.Status.VERIFIED, Achievement.Status.IQAC_LOCKED],
            achievement__is_deleted=False,
        )
        .select_related("achievement__faculty__department", "sub_criterion__criterion")
        .order_by("sub_criterion__code", "achievement__achievement_date")
    )

    grouped = {}
    for mapping in mappings:
        sub = mapping.sub_criterion
        grouped.setdefault(sub.criterion, {}).setdefault(sub, []).append(mapping.achievement)
    return grouped


def generate_report_docx(framework, cycle) -> tuple[bytes, list[str]]:
    """Returns (file_bytes, evidence_ids_included)."""
    grouped = _grouped_mappings(framework.id, cycle.id)

    doc = Document()
    title = doc.add_heading(f"{framework.name}: Accreditation Evidence Report", level=0)
    title.runs[0].font.size = Pt(22)

    meta = doc.add_paragraph()
    meta.add_run(f"Cycle: {cycle.label}\n").bold = True
    meta.add_run(f"Generated: {timezone.now():%d %B %Y, %H:%M} IST\n")
    meta.add_run("Includes verified and IQAC-locked evidence only.\n")

    evidence_ids: list[str] = []

    if not grouped:
        doc.add_paragraph("No verified evidence is currently mapped to this framework/cycle.")

    for criterion in sorted(grouped, key=lambda c: c.code):
        doc.add_heading(f"Criterion {criterion.code}: {criterion.title}", level=1)
        sub_map = grouped[criterion]
        for sub in sorted(sub_map, key=lambda s: s.code):
            doc.add_heading(f"{sub.code}: {sub.description}", level=2)
            table = doc.add_table(rows=1, cols=5)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            for i, label in enumerate(["Evidence ID", "Faculty", "Title", "Issuing Organization", "Date"]):
                hdr[i].text = label
            for achievement in sub_map[sub]:
                row = table.add_row().cells
                row[0].text = achievement.evidence_id
                row[1].text = achievement.faculty.full_name
                row[2].text = achievement.title
                row[3].text = achievement.issuing_organization or "Not stated"
                row[4].text = achievement.achievement_date.isoformat() if achievement.achievement_date else "Not stated"
                evidence_ids.append(achievement.evidence_id)
            doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    file_bytes = buf.getvalue()
    return file_bytes, sorted(set(evidence_ids))


def checksum_of(file_bytes: bytes) -> str:
    return hashlib.sha256(file_bytes).hexdigest()
