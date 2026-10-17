"""Best-effort text extraction (OCR) for uploaded evidence. Extracts text +
a confidence score from any supported attachment and guesses a few
metadata fields (date, issuing organization) to pre-fill the submission form.

Supported: JPG/PNG (Tesseract, with rotation fix + preprocessing), PDF
(real text first, OCR for scans), Word, PowerPoint, Excel and CSV (read
directly). Legacy .doc/.ppt/.xls and office files that are really just
pictures (a slide deck of screenshots, a scanned page pasted into Word) are
converted to PDF with LibreOffice and OCR'd.

This must NEVER block a submission: every call is wrapped so a failure
(missing binary, corrupt file, unreadable scan) degrades to "no text
extracted" rather than raising — Phase 4.4's hard requirement.
"""

import io
import logging
import re

logger = logging.getLogger("faems")

KNOWN_ORGANIZATIONS = [
    "SWAYAM", "NPTEL", "Coursera", "edX", "Udemy", "AWS", "Google", "Microsoft",
    "IEEE", "ACM", "Infosys", "TCS", "IBM", "NVIDIA", "Cisco", "Oracle",
]

DATE_PATTERNS = [
    r"\b(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})\b",
    r"\b(\d{4}-\d{2}-\d{2})\b",
    r"\b(\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s+\d{4})\b",
]

# Anything with less real text than this is treated as "probably pictures".
MIN_USEFUL_CHARS = 80
OCR_TIMEOUT = 40  # seconds per Tesseract call; this runs in a background thread


# ----------------------------------------------------------------- images

def _prepare(image):
    """Phone photos and scans: fix EXIF rotation, flatten transparency, make
    small images bigger (Tesseract reads ~300 DPI text best) and boost contrast."""
    from PIL import Image, ImageOps

    image = ImageOps.exif_transpose(image)
    if image.mode in ("RGBA", "LA", "P"):
        background = Image.new("RGB", image.size, "white")
        rgba = image.convert("RGBA")
        background.paste(rgba, mask=rgba.split()[-1])
        image = background
    else:
        image = image.convert("RGB")
    if image.width < 1600:
        scale = 1600 / image.width
        image = image.resize((int(image.width * scale), int(image.height * scale)), Image.LANCZOS)
    if image.width > 4000:
        scale = 4000 / image.width
        image = image.resize((int(image.width * scale), int(image.height * scale)), Image.LANCZOS)
    return ImageOps.autocontrast(ImageOps.grayscale(image))


def _ocr_pass(image, psm: int, timeout: int):
    import pytesseract

    data = pytesseract.image_to_data(
        image, output_type=pytesseract.Output.DICT, timeout=timeout, config=f"--psm {psm}"
    )
    words, confidences = [], []
    for conf, word in zip(data["conf"], data["text"]):
        if word.strip():
            words.append(word)
            try:
                if int(float(conf)) >= 0:
                    confidences.append(int(float(conf)))
            except ValueError:
                pass
    text = " ".join(words)
    avg = (sum(confidences) / len(confidences)) if confidences else 0.0
    return text, round(avg, 1)


def _ocr_image(image, timeout: int = OCR_TIMEOUT) -> tuple[str, float]:
    """Tries the common Tesseract page-segmentation modes and keeps the
    result with the most confident text: 3 = automatic, 6 = one uniform
    block (good for certificates), 11 = sparse text (good for posters and
    photos with scattered words). The surrounding try/except in the callers
    turns a Tesseract timeout (RuntimeError) into "no text", never a failure."""
    image = _prepare(image)
    best_text, best_conf = "", 0.0
    for psm in (3, 6, 11):
        text, conf = _ocr_pass(image, psm, timeout)
        score = conf * min(len(text), 400)
        if text and score > best_conf * min(len(best_text), 400):
            best_text, best_conf = text, conf
        # Good enough — don't spend more passes on a clean read.
        if len(text) >= 60 and conf >= 80:
            break
    # Sideways or upside-down pages (phone photos, scanner feeds): if the
    # read is poor, try the other three orientations and keep the best.
    if best_conf < 60:
        for angle in (90, 180, 270):
            text, conf = _ocr_pass(image.rotate(angle, expand=True, fillcolor=255), 3, timeout)
            if text and conf * min(len(text), 400) > best_conf * min(len(best_text), 400):
                best_text, best_conf = text, conf
    # Pictures without real text (scenery, people, a poster) make Tesseract
    # hallucinate gibberish at very low confidence. Treat that as "no text"
    # rather than feeding noise to the analysis.
    if best_conf < 45:
        return "", best_conf
    return best_text, best_conf


# -------------------------------------------------------------------- pdf

def _pdf_native_text(data: bytes) -> str:
    """Digital PDFs carry real text — far more accurate (and much faster)
    than OCR-ing a rendered image of it."""
    import subprocess

    try:
        result = subprocess.run(
            ["pdftotext", "-l", "8", "-layout", "-", "-"], input=data, capture_output=True, timeout=30,
        )
        return result.stdout.decode("utf-8", errors="ignore").strip()
    except Exception:  # noqa: BLE001
        return ""


def _pdf_text(data: bytes, max_pages: int = 4) -> tuple[str, float | None]:
    native = _pdf_native_text(data)
    if len(native) >= 40:
        return native, 100.0

    from pdf2image import convert_from_bytes

    pages = convert_from_bytes(data, first_page=1, last_page=max_pages, dpi=200)
    texts, confs = [], []
    for page in pages:
        t, c = _ocr_image(page)
        if t:
            texts.append(t)
            confs.append(c)
    if not texts:
        return "", None
    return "\n".join(texts), round(sum(confs) / len(confs), 1)


# ----------------------------------------------------------------- office

def _shape_text(shape, out):
    """Text from a PowerPoint shape: text frames, tables and grouped shapes."""
    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        out.append(shape.text_frame.text.strip())
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append(" | ".join(cells))
    if getattr(shape, "shapes", None) is not None:
        for sub in shape.shapes:
            _shape_text(sub, out)


def _office_text(file_obj, content_type: str) -> str:
    if "wordprocessingml" in content_type:
        import docx

        document = docx.Document(file_obj)
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = []
                for c in row.cells:
                    if c.text.strip() and c.text.strip() not in cells:
                        cells.append(c.text.strip())
                if cells:
                    parts.append(" | ".join(cells))
        for section in document.sections:
            for block in (section.header, section.footer):
                parts.extend(p.text for p in block.paragraphs if p.text.strip())
        return "\n".join(parts)
    if "presentationml" in content_type:
        from pptx import Presentation

        deck = Presentation(file_obj)
        parts = []
        for number, slide in enumerate(deck.slides, start=1):
            slide_parts = []
            for shape in slide.shapes:
                _shape_text(shape, slide_parts)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_parts.append(f"Notes: {notes}")
            if slide_parts:
                parts.append(f"Slide {number}: " + " ".join(slide_parts))
        return "\n".join(parts)
    if "spreadsheetml" in content_type:
        import openpyxl

        book = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        parts = []
        for sheet in book.worksheets[:6]:
            parts.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(max_row=300, max_col=30, values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if content_type in {"text/csv", "text/plain"}:
        return file_obj.read().decode("utf-8", errors="ignore")
    return ""


_LEGACY_OR_OFFICE = {
    "application/msword",
    "application/vnd.ms-powerpoint",
    "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _office_via_pdf(data: bytes, content_type: str) -> tuple[str, float | None]:
    """Old .doc/.ppt/.xls, and office files whose content is mostly pictures:
    render with LibreOffice and read the PDF (real text or OCR)."""
    from apps.achievements.preview import convert_to_pdf

    pdf = convert_to_pdf(data, content_type, timeout=120)
    if not pdf:
        return "", None
    return _pdf_text(pdf, max_pages=5)


# ------------------------------------------------------------------ public

def extract_text_and_confidence(file_obj, content_type: str) -> tuple[str, float | None]:
    """Returns (text, confidence_0_to_100). confidence is None if OCR
    wasn't attempted at all (e.g. unsupported type) rather than 0, so the
    caller can tell "low confidence" apart from "never tried". Digital
    documents (text PDFs, Word, Excel, PowerPoint, CSV) are read directly
    and reported with confidence 100."""
    try:
        file_obj.seek(0)
        if content_type in {"image/jpeg", "image/png"}:
            from PIL import Image

            return _ocr_image(Image.open(file_obj))

        if content_type == "application/pdf":
            return _pdf_text(file_obj.read())

        if content_type in _LEGACY_OR_OFFICE:
            data = file_obj.read()
            text, conf = "", None
            if content_type not in {
                "application/msword", "application/vnd.ms-powerpoint", "application/vnd.ms-excel",
            }:
                try:
                    text = _office_text(io.BytesIO(data), content_type)
                    conf = 100.0 if text.strip() else None
                except Exception as exc:  # noqa: BLE001 — fall through to the LibreOffice route
                    logger.info("Direct office read failed (%s); trying LibreOffice", exc)
            if len(text.strip()) < MIN_USEFUL_CHARS:
                rendered, rconf = _office_via_pdf(data, content_type)
                if len(rendered.strip()) > len(text.strip()):
                    text, conf = (text + "\n" + rendered).strip(), rconf or conf
            return text, conf

        text = _office_text(file_obj, content_type)
        return (text, 100.0) if text.strip() else ("", None)
    except Exception as exc:  # noqa: BLE001 — OCR must degrade, never block (Phase 4.4)
        logger.warning("OCR extraction failed, continuing without it: %s", exc)
        return "", None
    finally:
        try:
            file_obj.seek(0)
        except Exception:  # noqa: BLE001
            pass


def guess_fields(text: str) -> dict:
    """Cheap heuristic extraction — the faculty always confirms/edits these
    before save (Phase 4.3), so false positives here are cheap."""
    if not text:
        return {}

    fields = {}
    for org in KNOWN_ORGANIZATIONS:
        if re.search(rf"\b{re.escape(org)}\b", text, re.IGNORECASE):
            fields["issuing_organization"] = org
            break

    for pattern in DATE_PATTERNS:
        match = re.search(pattern, text)
        if match:
            fields["date_guess"] = match.group(1)
            break

    return fields
